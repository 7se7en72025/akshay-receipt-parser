import copy
import json
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import database as store
from .config import EXPORTS, TEMPLATES
from .extraction import clean
from .models import Invoice, TemplateConfig

FIELDS = {
    "bill_number": "Bill No.",
    "bill_date": "Bill Date",
    "supplier_name": "Supplier",
    "supplier_gstin": "Supplier GSTIN",
    "purchase_number": "Purchase No.",
    "description": "Product",
    "hsn": "HSN",
    "quantity": "Quantity",
    "unit": "Unit",
    "unit_price": "Rate",
    "discount_1": "Discount 1",
    "discount_2": "Discount 2",
    "taxable_amount": "Printed Item Amount",
    "sgst_rate": "SGST %",
    "sgst_amount": "SGST Amount",
    "cgst_rate": "CGST %",
    "cgst_amount": "CGST Amount",
    "igst_rate": "IGST %",
    "igst_amount": "IGST Amount",
    "cess_amount": "Cess Amount",
    "source_file": "Source File",
    "source_page": "Source Page",
    "currency": "Currency",
    "amounts_include_tax": "Item Amounts Include Tax",
    "secondary_amount": "Original Currency Amount",
    "secondary_currency": "Original Currency",
}
SUMMARY_FIELDS = {
    "bill_number": "Bill No.",
    "bill_date": "Bill Date",
    "supplier_name": "Supplier",
    "supplier_gstin": "Supplier GSTIN",
    "purchase_number": "Purchase No.",
    "taxable_total": "Taxable Total",
    "sgst_total": "SGST Total",
    "cgst_total": "CGST Total",
    "igst_total": "IGST Total",
    "cess_total": "Cess Total",
    "other_tax_total": "Other Tax / VAT",
    "charges_total": "Shipping / Other Charges",
    "discount_total": "Bill Discount",
    "withholding_total": "Withholding Tax",
    "amounts_include_tax": "Item Amounts Include Tax",
    "round_off": "Round Off",
    "grand_total": "Grand Total",
    "currency": "Currency",
    "secondary_subtotal": "Original Subtotal",
    "secondary_currency": "Original Currency",
    "exchange_rate": "Exchange Rate",
    "source_file": "Source File",
}
NUMERIC = {
    "secondary_amount",
    "quantity",
    "unit_price",
    "discount_1",
    "discount_2",
    "taxable_amount",
    "sgst_rate",
    "sgst_amount",
    "cgst_rate",
    "cgst_amount",
    "igst_rate",
    "igst_amount",
    "cess_amount",
}
ALIASES = {clean(label): key for key, label in FIELDS.items()}
ALIASES.update(
    {
        "qty": "quantity",
        "item": "description",
        "itemdescription": "description",
        "description": "description",
        "descriptionofgoods": "description",
        "invoiceno": "bill_number",
        "invoicenumber": "bill_number",
        "billnumber": "bill_number",
        "invoicedate": "bill_date",
        "suppliername": "supplier_name",
        "uom": "unit",
        "unitprice": "unit_price",
        "taxablevalue": "taxable_amount",
        "taxableamount": "taxable_amount",
        "cgst": "cgst_amount",
        "sgst": "sgst_amount",
        "igst": "igst_amount",
        "filename": "source_file",
    }
)


def check_xlsx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 5000:
                raise ValueError("Workbook contains too many files.")
            if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 40 MB.")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("Upload an Excel .xlsx workbook.")
    except zipfile.BadZipFile as error:
        raise ValueError("This file is not a readable .xlsx workbook.") from error


def inspect_template(path: Path, sheet_name=None, header_row=None):
    check_xlsx(path)
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = wb[sheet_name] if sheet_name else wb.active
        if header_row is None:
            candidates = list(sheet.iter_rows(min_row=1, max_row=40, max_col=60, values_only=True))
            scored = [
                (sum(clean(value) in ALIASES for value in row if value is not None), index + 1)
                for index, row in enumerate(candidates)
                if any(value is not None for value in row)
            ]
            if not scored:
                raise ValueError("No headings were found in the first 40 rows.")
            header_row = max(scored, key=lambda candidate: (candidate[0], -candidate[1]))[1]
        if not 1 <= header_row <= 200:
            raise ValueError("Choose a header row from 1 to 200.")
        values = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row, max_col=60, values_only=True),
            (),
        )
        columns = [
            {
                "index": index + 1,
                "letter": get_column_letter(index + 1),
                "label": str(value) if value is not None else "",
                "suggested_field": ALIASES.get(clean(value), ""),
            }
            for index, value in enumerate(values)
            if value is not None
        ]
        if not columns:
            raise ValueError("This row has no headings. Choose another header row.")
        mapping = {}
        for column in columns:
            if column["suggested_field"] and column["suggested_field"] not in mapping:
                mapping[column["suggested_field"]] = column["index"]
        # Suggested append row is informational; staff explicitly confirms it in the UI.
        return {
            "sheet_names": wb.sheetnames,
            "sheet_name": sheet.title,
            "header_row": header_row,
            "suggested_start_row": max(header_row + 1, sheet.max_row + 1),
            "columns": columns,
            "suggested_mapping": mapping,
        }
    finally:
        wb.close()


def assign(cell, value, numeric=False):
    if value is None:
        return
    if numeric:
        cell.value = Decimal(str(value))
        cell.number_format = "#,##0.00"
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = "dd-mm-yyyy"
    else:
        cell.value = str(value)
        cell.data_type = "s"  # Source strings must never become executable formulas.


def rows_for(documents):
    rows = []
    for document in documents:
        invoice = Invoice.model_validate(document["invoice"])
        common = {
            key: getattr(invoice, key)
            for key in [
                "bill_number",
                "bill_date",
                "supplier_name",
                "supplier_gstin",
                "purchase_number",
                "currency",
                "amounts_include_tax",
            ]
        }
        for item in invoice.items:
            rows.append({**common, **item.model_dump(), "source_file": document["filename"]})
    if len(rows) > 10000:
        raise ValueError("Export is limited to 10,000 product rows. Select fewer bills.")
    return rows


def style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E35")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for column in sheet.columns:
        index = column[0].column
        label = str(column[0].value or "")
        width = 48 if label in {"Product", "Source File"} else 24 if "Supplier" in label else 18
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.data_type == "n":
                cell.number_format = "#,##0.00"


def default_workbook():
    """The fixed built-in format; headings are shared with actual exports."""
    wb = Workbook()
    items = wb.active
    items.title = "Line Items"
    items.append(list(FIELDS.values()))
    summary = wb.create_sheet("Bill Summary")
    summary.append(list(SUMMARY_FIELDS.values()))
    style_sheet(items)
    style_sheet(summary)
    return wb


def create_export(documents, template=None):
    rows = rows_for(documents)
    if template:
        config = TemplateConfig.model_validate(json.loads(template["config_json"]))
        if not config.mapping or len(set(config.mapping.values())) != len(config.mapping):
            raise ValueError("Save a unique template column mapping first.")
        if any(
            key not in FIELDS or column < 1 or column > 60 for key, column in config.mapping.items()
        ):
            raise ValueError("Template mapping contains an unsupported field or column.")
        if config.start_row <= config.header_row:
            raise ValueError("The insertion row must be after the header row.")
        path = TEMPLATES / template["storage_name"]
        check_xlsx(path)
        wb = load_workbook(path, data_only=False, keep_links=False)
        sheet = wb[config.sheet_name]
        for offset, row in enumerate(rows):
            for field, column in config.mapping.items():
                cell = sheet.cell(config.start_row + offset, column)
                merged = isinstance(cell, MergedCell) or any(
                    cell.coordinate in area for area in sheet.merged_cells.ranges
                )
                if merged or cell.value is not None:
                    wb.close()
                    raise ValueError(
                        f"Template cell {cell.coordinate} is occupied or merged. "
                        "Choose an empty insertion row; existing data is never overwritten."
                    )
                if offset:
                    seed = sheet.cell(config.start_row, column)
                    cell._style = copy.copy(seed._style)
                assign(cell, row.get(field), field in NUMERIC)
    else:
        wb = default_workbook()
        sheet = wb["Line Items"]
        for row_index, row in enumerate(rows, start=2):
            for column, field in enumerate(FIELDS, start=1):
                assign(sheet.cell(row_index, column), row.get(field), field in NUMERIC)
        style_sheet(sheet)
        summary = wb["Bill Summary"]
        for row_index, document in enumerate(documents, start=2):
            invoice = Invoice.model_validate(document["invoice"])
            for column, field in enumerate(SUMMARY_FIELDS, start=1):
                value = document["filename"] if field == "source_file" else getattr(invoice, field)
                assign(
                    summary.cell(row_index, column),
                    value,
                    field.endswith("_total")
                    or field in {"round_off", "secondary_subtotal", "exchange_rate"},
                )
        style_sheet(summary)
    export_id = uuid4().hex
    storage_name = export_id + ".xlsx"
    try:
        wb.save(EXPORTS / storage_name)
    finally:
        wb.close()
    snapshot = [{"id": doc["id"], "version": doc["version"]} for doc in documents]
    with store.connection() as db:
        db.execute(
            "INSERT INTO exports VALUES(?,?,?,?)",
            (
                export_id,
                storage_name,
                json.dumps(
                    {
                        "documents": snapshot,
                        "template_id": template["id"] if template else None,
                        "template_config": json.loads(template["config_json"])
                        if template
                        else None,
                        "rows": len(rows),
                    }
                ),
                store.now(),
            ),
        )
        for document in documents:
            store.event(db, document["id"], "exported", {"export_id": export_id})
    return {"id": export_id, "row_count": len(rows), "bill_count": len(documents)}
