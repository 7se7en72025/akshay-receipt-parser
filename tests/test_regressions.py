import copy
import io
import itertools
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import database as store
from backend.ai import merge_proposal
from backend.extraction import header_fields, table_items
from backend.layouts import generic_fields
from backend.models import Invoice, Item, number
from backend.validation import validate_invoice


def valid_invoice():
    return Invoice(
        supplier_name="Example Supplier",
        bill_number="TEST-001",
        bill_date=date(2026, 1, 2),
        taxable_total=100,
        cgst_total=9,
        sgst_total=9,
        igst_total=0,
        cess_total=0,
        grand_total=118,
        items=[
            Item(
                description="Source widget",
                quantity=2,
                unit="Pc.",
                unit_price=50,
                taxable_amount=100,
                cgst_amount=9,
                sgst_amount=9,
                igst_amount=0,
                cess_amount=0,
            )
        ],
    )


class ParsingRegressions(unittest.TestCase):
    def test_all_120_column_orders(self):
        fields = {
            "Description": "Source widget",
            "Qty": "2",
            "Unit Price": "50",
            "Amount": "100",
            "Unit": "Pc.",
        }
        for headers in itertools.permutations(fields):
            with self.subTest(headers=headers):
                item = table_items([list(headers), [fields[key] for key in headers]], 1)[0]
                self.assertEqual(
                    (
                        item.description,
                        item.quantity,
                        item.unit_price,
                        item.taxable_amount,
                        item.unit,
                    ),
                    ("Source widget", Decimal(2), Decimal(50), Decimal(100), "Pc."),
                )

    def test_space_grouped_quantity_not_misread_as_unit(self):
        item = table_items(
            [
                ["Description", "Qty / Value", "Rate", "Amount"],
                ["Source fee", "23 026.230", "18.00%", "4144.72"],
            ],
            1,
        )[0]
        self.assertEqual(item.quantity, Decimal("23026.230"))
        self.assertEqual(item.unit, "")

    def test_regional_tax_and_foreign_registration(self):
        data = generic_fields(
            "Seller INVOICE\nTVA: FR240529481\nSubtotal $1108\nOH State Sales Tax (5.75%) $63.71\nFranklin County Tax (1.75%) $19.39\nTOTAL $1191.10"
        )
        self.assertEqual(data["other_tax_total"], Decimal("83.10"))

    def test_locale_numbers(self):
        for source, expected in [
            ("10,514.50", "10514.50"),
            ("1.151,18", "1151.18"),
            ("1 857,44", "1857.44"),
            ("INR 12,34,567.89", "1234567.89"),
            ("0,00", "0"),
            ("18.45%", "18.45"),
            ("1,000", "1000"),
        ]:
            with self.subTest(source=source):
                self.assertEqual(number(source), Decimal(expected))

    def test_invalid_number_not_invented(self):
        for value in ("unreadable", "12/34", "1O.50", "", None):
            self.assertIsNone(number(value))

    def test_footer_lines_do_not_join(self):
        self.assertEqual(
            header_fields(
                "Total Taxable Value 53651.35\n53651.35 2.50%\nGrand Total (INR) 56,334.00"
            )["taxable_total"],
            Decimal("53651.35"),
        )

    def test_invoice_title_not_invoice_identifier(self):
        self.assertNotIn(
            "bill_number", header_fields("INVOICE Northwind Logistics\nDate: 03/12/2024")
        )
        self.assertEqual(
            generic_fields(
                "INVOICE Northwind Logistics\nInvoice #: INV-123\nDate: 03/12/2024\nColumbus, OH 43215"
            )["bill_number"],
            "INV-123",
        )

    def test_tax_registration_is_not_tax_amount(self):
        data = generic_fields(
            "Seller INVOICE\nVAT Reg No: 4349344539\nBILL TO\nVAT No: 4140229006\nSubtotal N$ 100.00\nVAT 15% N$15.00\nTOTAL N$115.00"
        )
        self.assertEqual(data["other_tax_total"], Decimal(15))

    def test_named_and_year_first_dates(self):
        for source, expected in [
            ("July 30, 2026", date(2026, 7, 30)),
            ("8 August 2025", date(2025, 8, 8)),
            ("2026/06/19", date(2026, 6, 19)),
        ]:
            with self.subTest(source=source):
                self.assertEqual(
                    generic_fields(f"Seller INVOICE\nDate: {source}\nBILL TO")["bill_date"],
                    expected,
                )

    def test_reordered_columns(self):
        rows = [
            ["Amount", "Description", "Unit Price", "Qty", "Unit"],
            ["100.00", "Source widget", "50.00", "2", "Pc."],
        ]
        item = table_items(rows, 1)[0]
        self.assertEqual(
            (item.description, item.quantity, item.unit_price, item.taxable_amount),
            ("Source widget", Decimal(2), Decimal(50), Decimal(100)),
        )

    def test_one_invalid_field_keeps_valid_fields(self):
        item = table_items(
            [
                ["Description", "Qty", "Rate", "Amount"],
                ["Source widget", "2", "999999999999999999999", "100"],
            ],
            1,
        )[0]
        self.assertEqual(item.quantity, Decimal(2))
        self.assertEqual(item.taxable_amount, Decimal(100))
        self.assertIsNone(item.unit_price)

    def test_dual_amount_columns_keep_both(self):
        item = table_items(
            [
                ["Description", "Qty", "Amount USD", "Amount ZAR"],
                ["Source widget", "1", "27.56", "508.52"],
            ],
            1,
        )[0]
        self.assertEqual(item.taxable_amount, Decimal("508.52"))
        self.assertEqual(item.secondary_amount, Decimal("27.56"))
        self.assertEqual(item.secondary_currency, "USD")

    def test_inclusive_amounts_are_not_rewritten(self):
        data = generic_fields(
            "Seller INVOICE\nSubtotal N$115.00\nVAT 15% (included) N$15.00\nTOTAL N$115.00"
        )
        self.assertEqual(data["taxable_total"], Decimal(100))
        self.assertTrue(data["amounts_include_tax"])

    def test_total_including_vat_is_not_added_as_a_second_tax(self):
        data = generic_fields(
            "Seller INVOICE\nTOTAL (VAT inclusive) N$115.00\nVAT 15% included in total N$15.00"
        )
        self.assertEqual(data["other_tax_total"], Decimal(15))
        self.assertEqual(data["taxable_total"], Decimal(100))


class PreservationRegressions(unittest.TestCase):
    def test_validation_never_mutates_entries(self):
        invoice = valid_invoice()
        invoice.grand_total = Decimal(999)
        before = invoice.model_dump_json()
        result = validate_invoice(invoice)
        self.assertGreater(result["errors"], 0)
        self.assertEqual(invoice.model_dump_json(), before)

    def test_correct_lines_with_wrong_total_are_retained(self):
        invoice = valid_invoice()
        invoice.taxable_total = Decimal(150)
        original_items = copy.deepcopy(invoice.items)
        self.assertFalse(validate_invoice(invoice)["can_approve"])
        self.assertEqual(invoice.items, original_items)

    def test_blank_and_foreign_bills_block_approval(self):
        invoice = valid_invoice()
        invoice.supplier_name = "   "
        self.assertFalse(validate_invoice(invoice)["can_approve"])
        invoice.supplier_name = "Supplier"
        invoice.currency = "USD"
        self.assertTrue(
            any(issue["field"] == "currency" for issue in validate_invoice(invoice)["issues"])
        )

    def test_fx_one_cent_difference_is_flagged(self):
        invoice = valid_invoice()
        invoice.currency = "NAD"
        invoice.secondary_subtotal = Decimal("162.48")
        invoice.exchange_rate = Decimal("19.9871")
        invoice.grand_total = Decimal("3247.51")
        before = invoice.model_dump_json()
        issues = validate_invoice(invoice)["issues"]
        self.assertTrue(any(issue["field"] == "exchange_rate" for issue in issues))
        self.assertEqual(invoice.model_dump_json(), before)

    def test_ai_conflicts_do_not_overwrite_correct_source(self):
        original = valid_invoice()
        proposal = original.model_copy(deep=True)
        proposal.grand_total = Decimal(999)
        proposal.items[0].quantity = Decimal(20)
        proposal.items[0].taxable_amount = Decimal(1000)
        before = original.model_dump_json()
        merged, notes = merge_proposal(
            original, proposal, "Source widget 2 Pc. 50.00 100.00 118.00"
        )
        self.assertEqual(merged.model_dump_json(), before)
        self.assertEqual(original.model_dump_json(), before)
        self.assertTrue(notes)

    def test_ai_can_fill_source_backed_gap(self):
        original = valid_invoice()
        original.items[0].unit_price = None
        proposal = valid_invoice()
        merged, _ = merge_proposal(original, proposal, "Source widget 2 Pc. 50.00 100.00")
        self.assertEqual(merged.items[0].unit_price, Decimal(50))
        self.assertIsNone(original.items[0].unit_price)

    def test_ai_rejects_unbacked_numbers(self):
        original = valid_invoice()
        original.items[0].unit_price = None
        proposal = valid_invoice()
        proposal.items[0].unit_price = Decimal(777)
        merged, _ = merge_proposal(original, proposal, "Source widget 2 Pc. 100.00")
        self.assertIsNone(merged.items[0].unit_price)

    def test_ai_row_count_cannot_drop_source_rows(self):
        original = valid_invoice()
        proposal = original.model_copy(deep=True)
        proposal.items = []
        merged, _ = merge_proposal(original, proposal, "Source widget")
        self.assertEqual(merged.items, original.items)


class ApiWorkflowRegressions(unittest.TestCase):
    def test_default_format_is_fixed_and_matches_export(self):
        from backend.workbooks import FIELDS, SUMMARY_FIELDS

        response = self.client.get("/api/templates/default/download")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        try:
            self.assertEqual(workbook.sheetnames, ["Line Items", "Bill Summary"])
            self.assertEqual(
                [cell.value for cell in workbook["Line Items"][1]], list(FIELDS.values())
            )
            self.assertEqual(
                [cell.value for cell in workbook["Bill Summary"][1]], list(SUMMARY_FIELDS.values())
            )
            self.assertTrue(all(sheet.max_row == 1 for sheet in workbook))
        finally:
            workbook.close()
        self.assertEqual(self.client.get("/api/templates").json(), [])
        self.assertEqual(self.client.get("/api/exports").json(), [])
        self.assertEqual(
            self.client.put(
                "/api/templates/default",
                json={"sheet_name": "Line Items", "mapping": {"description": 1}},
            ).status_code,
            404,
        )

    def test_template_merged_anchor_cannot_be_overwritten(self):
        import json

        from openpyxl import Workbook

        from backend.workbooks import create_export

        wb = Workbook()
        wb.active.title = "Items"
        wb.active["A1"] = "Product"
        wb.active.merge_cells("A2:B2")
        path = self.root / "templates" / "merged.xlsx"
        wb.save(path)
        wb.close()
        original = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "occupied or merged"):
            create_export(
                [
                    {
                        "id": "bill",
                        "version": 1,
                        "filename": "source.pdf",
                        "invoice": self.invoice.model_dump(mode="json"),
                    }
                ],
                {
                    "id": "merged",
                    "storage_name": "merged.xlsx",
                    "config_json": json.dumps(
                        {
                            "sheet_name": "Items",
                            "header_row": 1,
                            "start_row": 2,
                            "mapping": {"description": 1},
                        }
                    ),
                },
            )
        self.assertEqual(path.read_bytes(), original)

    def setUp(self):
        import backend.main as main
        import backend.workbooks as workbooks

        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.patches = [
            patch.object(store, "DATA", self.root),
            patch.object(main, "UPLOADS", self.root / "uploads"),
            patch.object(main, "EXPORTS", self.root / "exports"),
            patch.object(main, "TEMPLATES", self.root / "templates"),
            patch.object(workbooks, "EXPORTS", self.root / "exports"),
            patch.object(workbooks, "TEMPLATES", self.root / "templates"),
        ]
        for folder in ("uploads", "exports", "templates"):
            (self.root / folder).mkdir()
        for item in self.patches:
            item.start()
        self.client = TestClient(main.app, base_url="http://localhost")
        self.client.__enter__()
        self.invoice = valid_invoice()
        with store.connection() as db:
            db.execute(
                """INSERT INTO documents(id,filename,storage_name,mime,sha256,status,invoice_json,created_at,updated_at,source_currency)
                VALUES('bill','source.pdf','source.pdf','application/pdf','unique','needs_review',?,?,?,'INR')""",
                (self.invoice.model_dump_json(), store.now(), store.now()),
            )

    def tearDown(self):
        self.client.__exit__(None, None, None)
        for item in reversed(self.patches):
            item.stop()
        self.temporary.cleanup()

    def test_wrong_total_does_not_change_saved_entries(self):
        invoice = self.invoice.model_dump(mode="json")
        invoice["grand_total"] = "999"
        result = self.client.post(
            "/api/documents/bill/approve", json={"version": 1, "invoice": invoice}
        )
        self.assertEqual(result.status_code, 422)
        saved = self.client.get("/api/documents/bill").json()
        self.assertEqual(saved["invoice"]["grand_total"], "118")
        self.assertEqual(saved["invoice"]["items"], self.invoice.model_dump(mode="json")["items"])

    def test_save_approve_export_reopen_and_restore(self):
        payload = {"version": 1, "invoice": self.invoice.model_dump(mode="json")}
        saved = self.client.put("/api/documents/bill", json=payload).json()
        event = saved["history"][0]["id"]
        approved = self.client.post(
            "/api/documents/bill/approve", json={**payload, "version": saved["version"]}
        )
        self.assertEqual(approved.status_code, 200)
        export = self.client.post("/api/exports", json={"document_ids": ["bill"]})
        self.assertEqual(export.status_code, 201)
        download = self.client.get(f"/api/exports/{export.json()['id']}/download")
        self.assertEqual(download.status_code, 200)
        restored = self.client.post(
            f"/api/documents/bill/history/{event}/restore",
            json={"version": approved.json()["version"]},
        )
        self.assertEqual(restored.status_code, 200)
        self.assertEqual(restored.json()["status"], "needs_review")
        self.assertEqual(
            self.client.post("/api/exports", json={"document_ids": ["bill"]}).status_code, 422
        )

    def test_stale_version_and_currency_relabelling_rejected(self):
        payload = {"version": 1, "invoice": self.invoice.model_dump(mode="json")}
        self.assertEqual(self.client.put("/api/documents/bill", json=payload).status_code, 200)
        self.assertEqual(self.client.put("/api/documents/bill", json=payload).status_code, 409)
        payload["version"] = 2
        payload["invoice"]["currency"] = "USD"
        self.assertEqual(self.client.put("/api/documents/bill", json=payload).status_code, 422)

    def test_inbox_search_date_and_paging(self):
        for query in (
            "search=Source",
            "search=TEST-001",
            "date_from=2026-01-01&date_to=2026-01-31",
            "page=100",
        ):
            response = self.client.get("/api/inbox?" + query)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["matched"], 1)
            self.assertNotIn("items", response.json()["items"][0]["invoice"])
        self.assertEqual(
            self.client.get("/api/inbox?date_from=2026-02-01&date_to=2026-01-01").status_code, 422
        )
        self.assertEqual(self.client.get("/api/inbox?search=%25").json()["matched"], 0)

    def test_excel_formula_injection_is_stored_as_text(self):
        from backend.workbooks import create_export

        invoice = self.invoice.model_copy(deep=True)
        invoice.items[0].description = '=HYPERLINK("https://invalid.test","click")'
        result = create_export(
            [
                {
                    "id": "bill",
                    "version": 1,
                    "filename": "source.pdf",
                    "invoice": invoice.model_dump(mode="json"),
                }
            ]
        )
        wb = load_workbook(self.root / "exports" / (result["id"] + ".xlsx"))
        try:
            values = [
                cell
                for row in wb["Line Items"]
                for cell in row
                if cell.value == invoice.items[0].description
            ]
            self.assertEqual(len(values), 1)
            self.assertEqual(values[0].data_type, "s")
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
